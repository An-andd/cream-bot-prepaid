"""Optional .docx output using the same 2 x 3 bordered-table layout.

The PDF is the primary deliverable (that is what WhatsApp sends back), but the
same label content can be written into a Word document when you want to tweak
something by hand before printing.
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from . import config
from .labels import build_from_lines, build_label_lines, build_to_lines
from .parsing import Order


def _set_cell_borders(cell, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tc_pr.append(borders)


def _write_lines(cell, lines, font_size: float) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    first = True
    for text, bold in lines:
        if not first:
            paragraph = cell.add_paragraph()
        first = False
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.name = "Arial"


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)


def _fill_cell(cell, order: Order, biller_id: str, font_size: float) -> None:
    """Customer block left, sender block bottom-right, like the printed label."""
    cell.text = ""
    if config.FROM_BLOCK_POSITION != "right":
        _write_lines(cell, build_label_lines(order, biller_id), font_size)
        return

    inner = cell.add_table(rows=1, cols=2)
    _no_borders(inner)
    left, right = inner.cell(0, 0), inner.cell(0, 1)
    left.width = Inches(2.0)
    right.width = Inches(1.5)
    right.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    _write_lines(left, build_to_lines(order), font_size)
    _write_lines(right, build_from_lines(biller_id), font_size)
    # drop the empty paragraph Word leaves above a nested table
    if cell.paragraphs and not cell.paragraphs[0].text:
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)


def render_docx(orders: list[Order], biller_id: str, out_path: str,
                font_size: float = 9.0) -> str:
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc = Document(config.DOCX_TEMPLATE) if config.DOCX_TEMPLATE else Document()
    if config.DOCX_TEMPLATE:
        # Template is the master design: keep styles, drop any sample content.
        for element in list(doc.element.body):
            doc.element.body.remove(element)

    per_page = config.LABELS_PER_PAGE
    cols = config.LABELS_PER_ROW
    rows = config.LABEL_ROWS

    for page_start in range(0, max(len(orders), 1), per_page):
        chunk = orders[page_start:page_start + per_page]
        table = doc.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for idx in range(per_page):
            r, c = divmod(idx, cols)
            cell = table.cell(r, c)
            _set_cell_borders(cell)
            if idx < len(chunk):
                _fill_cell(cell, chunk[idx], biller_id, font_size)
        if page_start + per_page < len(orders):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.save(out_path)
    return out_path
